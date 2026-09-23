#!/usr/bin/env python3
"""
profile_workbook.py -- read-only structural profile of an .xlsx workbook.

Purpose
-------
This script bakes in NO assumptions about layout. It reports what is actually
stored in the file so that a parser can be written against reality rather than
against a print rendering (which hides merged cells, hidden rows/columns, real
sheet boundaries, formulas-vs-values, and true column positions).

It never writes to the workbook. It only reads.

Usage
-----
    pip install openpyxl
    python tools/profile_workbook.py "May sushi sales.xlsx"

    # options
    --rows N          rows of raw cell dump per sheet            (default 20)
    --max-cols N      column cap for the raw dump                (default 60)
    --scan-rows N     rows scanned for type/formula census       (default 5000)
    --out PATH        full report file    (default workbook_profile.txt)
    --json PATH       also emit a machine-readable JSON profile  (optional)
    --full            disable sheet grouping; print every sheet in full
    --reps N          detail dumps printed per layout group      (default 2)

Output
------
The FULL report always goes to --out. STDOUT gets a condensed version in which
sheets that share an identical layout fingerprint are collapsed to a few
representatives, so the output stays pasteable. Every sheet still appears in
the one-line SHEET INDEX regardless.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import os
import re
import sys
from collections import Counter, OrderedDict, defaultdict

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("ERROR: openpyxl is required.  pip install openpyxl")

try:
    from openpyxl.styles.numbers import is_date_format as _is_date_format
except Exception:  # pragma: no cover - older/newer openpyxl
    def _is_date_format(fmt):
        return bool(fmt) and re.search(r"[dmyhs]", str(fmt), re.I) is not None


MAX_REPR = 90
CELL_REF_RE = re.compile(r"\$?[A-Za-z]{1,3}\$?\d{1,7}")
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #

def short(obj, limit=MAX_REPR):
    """repr() truncated -- shows quoting, whitespace and unicode exactly."""
    try:
        text = repr(obj)
    except Exception as exc:
        text = "<unreprable %s: %s>" % (type(obj).__name__, exc)
    if len(text) > limit:
        text = text[: limit - 3] + "..."
    return text


def safe(fn, default=None):
    try:
        return fn()
    except Exception as exc:
        return default if default is not None else "<error: %s>" % exc


def type_label(cell):
    """A coarse, honest label for what is stored in a cell."""
    if cell.data_type == "f":
        return "formula"
    value = cell.value
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return "datetime"
    if isinstance(value, (int, float)):
        if safe(lambda: _is_date_format(cell.number_format), False):
            return "number(date-formatted)"
        return "number"
    if isinstance(value, str):
        if value == "":
            return "empty-str"
        if value.strip() == "":
            return "whitespace-str"
        if value.startswith("="):
            return "str-starting-with-eq"
        return "str"
    return type(value).__name__


def is_empty(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def formula_text(value):
    if isinstance(value, str):
        return value
    # ArrayFormula / shared formula objects
    return safe(lambda: str(getattr(value, "text", value)), short(value))


def formula_shape(text):
    """Normalise a formula so structurally identical ones collapse together."""
    shape = CELL_REF_RE.sub("<ref>", text)
    shape = re.sub(r"\d+(?:\.\d+)?", "<n>", shape)
    return shape


def sheets_referenced(text):
    out = set()
    for quoted, bare in SHEET_REF_RE.findall(text):
        name = (quoted or bare).replace("''", "'")
        if name and not re.fullmatch(r"[A-Za-z]{1,3}", name):
            out.add(name)
    return out


# --------------------------------------------------------------------------- #
# workbook-level
# --------------------------------------------------------------------------- #

def profile_workbook_meta(path, wb_f, wb_v, out, doc):
    size = os.path.getsize(path)
    with open(path, "rb") as fh:
        digest = hashlib.sha256(fh.read()).hexdigest()

    out("=" * 78)
    out("FILE")
    out("=" * 78)
    out("path            : %s" % os.path.abspath(path))
    out("size            : %s bytes" % size)
    out("sha256          : %s" % digest)
    out("mtime           : %s" % _dt.datetime.fromtimestamp(os.path.getmtime(path)))
    out("openpyxl        : %s" % openpyxl.__version__)
    out("python          : %s" % sys.version.split()[0])
    doc["file"] = {"path": os.path.abspath(path), "size": size, "sha256": digest,
                   "openpyxl": openpyxl.__version__}

    props = safe(lambda: wb_f.properties)
    out("")
    out("workbook properties")
    for attr in ("creator", "lastModifiedBy", "created", "modified", "title"):
        out("  %-14s: %s" % (attr, safe(lambda a=attr: getattr(props, a, None))))

    out("")
    out("=" * 78)
    out("WORKBOOK")
    out("=" * 78)
    out("sheet count     : %d" % len(wb_f.worksheets))
    out("sheet order (index, name, state):")
    sheets_meta = []
    for i, ws in enumerate(wb_f.worksheets):
        state = safe(lambda w=ws: w.sheet_state, "?")
        out("  [%2d] %-34s %s" % (i, short(ws.title, 34), state))
        sheets_meta.append({"index": i, "name": ws.title, "state": state})
    doc["sheets_order"] = sheets_meta

    # chartsheets / other non-worksheet sheets
    extra = [n for n in wb_f.sheetnames if n not in [w.title for w in wb_f.worksheets]]
    if extra:
        out("")
        out("NON-WORKSHEET sheets (chartsheets etc.): %s" % extra)

    # defined names
    out("")
    out("defined names:")
    items = []
    try:
        items = list(wb_f.defined_names.items())          # openpyxl >= 3.1
    except AttributeError:
        items = [(d.name, d) for d in safe(lambda: wb_f.defined_names.definedName, [])]
    except Exception:
        items = []
    for name, dn in items:
        val = safe(lambda d=dn: d.value, "?")
        out("  %-30s -> %s" % (short(name, 30), short(val, 120)))
    if not items:
        out("  (none)")
    doc["defined_names"] = [{"name": n, "value": str(safe(lambda d=d: d.value, "?"))}
                            for n, d in items]

    # external links
    ext = safe(lambda: list(wb_f._external_links), [])
    out("")
    out("external links  : %d" % len(ext))
    for link in ext:
        out("  %s" % short(safe(lambda l=link: l.file_link.Target, "?"), 140))

    # cached values present?  (decides whether formula results are readable)
    total_f = total_cached = 0
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    total_f += 1
                    if ws_v.cell(cell.row, cell.column).value is not None:
                        total_cached += 1
    out("")
    out("formula cells   : %d" % total_f)
    out("with cached val : %d" % total_cached)
    if total_f and total_cached == 0:
        out("  !! WARNING: no cached formula results. data_only=True yields None for")
        out("     every formula. The file must be opened+saved in Excel/LibreOffice,")
        out("     or every derived value must be recomputed by us.")
    doc["formula_cells"] = total_f
    doc["formula_cells_with_cached_value"] = total_cached


def profile_string_hygiene(wb_f, args, out, doc):
    """Text values that differ only by whitespace/case/unicode are the classic
    cause of an item silently splitting into two items during import. Report
    them as observations; do not guess which spelling is canonical."""
    raw_counts = Counter()
    where = {}
    for ws in wb_f.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, args.scan_rows),
                                max_col=min(ws.max_column, args.max_cols)):
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.strip() and cell.data_type != "f":
                    raw_counts[value] += 1
                    where.setdefault(value, "%s!%s" % (ws.title, cell.coordinate))

    normalised = defaultdict(set)
    for value in raw_counts:
        normalised[re.sub(r"\s+", " ", value.strip().lower())].add(value)
    collisions = {k: sorted(v) for k, v in normalised.items() if len(v) > 1}

    suspicious = sorted(
        v for v in raw_counts
        if v != v.strip() or "\xa0" in v or any(ord(ch) > 126 for ch in v))

    out("")
    out("=" * 78)
    out("STRING HYGIENE  (whitespace / case / unicode variants)")
    out("=" * 78)
    out("distinct text values : %d" % len(raw_counts))
    out("")
    out("values that collapse to the same normalised key (%d group(s)):" % len(collisions))
    for key, variants in sorted(collisions.items())[:60]:
        out("  %-38s %s" % (short(key, 38),
                            " | ".join("%s x%d" % (short(v, 32), raw_counts[v])
                                       for v in variants)))
    if len(collisions) > 60:
        out("  ... %d more groups" % (len(collisions) - 60))
    if not collisions:
        out("  (none)")
    out("")
    out("values with leading/trailing whitespace or non-ASCII characters (%d):"
        % len(suspicious))
    for value in suspicious[:60]:
        out("  %-46s x%-4d first seen %s" % (short(value, 46), raw_counts[value],
                                             where[value]))
    if len(suspicious) > 60:
        out("  ... %d more" % (len(suspicious) - 60))
    if not suspicious:
        out("  (none)")

    doc["string_collisions"] = collisions
    doc["string_suspicious"] = suspicious


# --------------------------------------------------------------------------- #
# sheet-level
# --------------------------------------------------------------------------- #

def sheet_extent(ws_f, scan_rows, max_cols):
    """True used extent, plus blank/ragged census. openpyxl's max_row/max_column
    over-report when formatting has been applied to empty cells."""
    last_row = 0
    last_col = 0
    blank_rows = 0
    blank_row_numbers = []
    row_widths = Counter()
    for row in ws_f.iter_rows(min_row=1, max_row=min(ws_f.max_row, scan_rows),
                              max_col=min(ws_f.max_column, max_cols)):
        width = 0
        for cell in row:
            if not is_empty(cell.value):
                width = max(width, cell.column)
        if width == 0:
            blank_rows += 1
            if len(blank_row_numbers) < 60 and row:
                blank_row_numbers.append(row[0].row)
        else:
            last_row = row[0].row
            last_col = max(last_col, width)
            row_widths[width] += 1
    return last_row, last_col, blank_rows, blank_row_numbers, row_widths


def profile_sheet(ws_f, ws_v, args, out, doc_sheet):
    name = ws_f.title
    out("")
    out("-" * 78)
    out("SHEET DETAIL: %s" % short(name))
    out("-" * 78)

    declared_rows = ws_f.max_row
    declared_cols = ws_f.max_column
    last_row, last_col, blank_rows, blank_row_nums, row_widths = sheet_extent(
        ws_f, args.scan_rows, args.max_cols)

    out("state                 : %s" % safe(lambda: ws_f.sheet_state, "?"))
    out("dimensions (declared) : %s   max_row=%s max_col=%s (%s)"
        % (safe(lambda: ws_f.dimensions, "?"), declared_rows, declared_cols,
           get_column_letter(max(1, min(declared_cols, 16384)))))
    out("used extent (scanned) : last non-empty row=%s col=%s (%s)"
        % (last_row, last_col, get_column_letter(last_col) if last_col else "-"))
    if declared_rows > last_row or declared_cols > last_col:
        out("  note: declared extent exceeds used extent -> formatting on empty cells")
    if declared_rows > args.scan_rows:
        out("  note: scan capped at --scan-rows=%d" % args.scan_rows)
    if declared_cols > args.max_cols:
        out("  note: columns capped at --max-cols=%d" % args.max_cols)

    out("freeze_panes          : %s" % safe(lambda: ws_f.freeze_panes))
    out("auto_filter           : %s" % safe(lambda: ws_f.auto_filter.ref))
    out("print_area            : %s" % safe(lambda: ws_f.print_area))
    out("charts / images       : %s / %s"
        % (len(safe(lambda: ws_f._charts, [])), len(safe(lambda: ws_f._images, []))))
    out("conditional formats   : %s" % len(list(safe(lambda: ws_f.conditional_formatting, []))))
    out("data validations      : %s" % len(safe(lambda: ws_f.data_validations.dataValidation, [])))

    tables = safe(lambda: dict(ws_f.tables), {})
    out("excel tables          : %s"
        % (", ".join("%s=%s" % (k, v.ref) for k, v in tables.items()) or "(none)"))

    rb = safe(lambda: [b.id for b in ws_f.row_breaks.brk], [])
    cb = safe(lambda: [b.id for b in ws_f.col_breaks.brk], [])
    out("manual page breaks    : rows=%s cols=%s" % (rb or "(none)", cb or "(none)"))

    # --- merged cells -------------------------------------------------------
    merged = sorted(str(r) for r in safe(lambda: ws_f.merged_cells.ranges, []))
    out("")
    out("MERGED RANGES (%d):" % len(merged))
    for chunk_start in range(0, len(merged), 8):
        out("  " + "  ".join(merged[chunk_start:chunk_start + 8]))
    if not merged:
        out("  (none)")

    # --- hidden rows / cols -------------------------------------------------
    hidden_rows = sorted(r for r, dim in safe(lambda: list(ws_f.row_dimensions.items()), [])
                         if getattr(dim, "hidden", False) or getattr(dim, "height", None) == 0)
    hidden_cols = []
    for letter, dim in safe(lambda: list(ws_f.column_dimensions.items()), []):
        if getattr(dim, "hidden", False) or getattr(dim, "width", None) == 0:
            lo, hi = getattr(dim, "min", None), getattr(dim, "max", None)
            if lo and hi:
                hidden_cols.extend(get_column_letter(c) for c in range(lo, hi + 1))
            else:
                hidden_cols.append(letter)
    hidden_cols = sorted(set(hidden_cols))
    out("")
    out("HIDDEN ROWS (%d): %s" % (len(hidden_rows), hidden_rows[:80] or "(none)"))
    out("HIDDEN COLS (%d): %s" % (len(hidden_cols), hidden_cols[:80] or "(none)"))
    out("  (a hidden column is invisible in any print rendering -- these are the")
    out("   columns most likely to be missing from the PDF)")

    # --- outline groups (another way columns disappear) ---------------------
    grouped = [(letter, getattr(d, "outlineLevel", 0))
               for letter, d in safe(lambda: list(ws_f.column_dimensions.items()), [])
               if getattr(d, "outlineLevel", 0)]
    if grouped:
        out("OUTLINE-GROUPED COLS  : %s" % grouped[:40])

    # --- raw cell dump ------------------------------------------------------
    out("")
    out("FIRST %d ROWS, EXACTLY AS STORED  (blank cells omitted)" % args.rows)
    out("  format: CELL | type | value_repr | fmt | [formula -> cached value]")
    dump_rows = min(args.rows, declared_rows)
    ncols = max(1, min(declared_cols, args.max_cols))
    for r in range(1, dump_rows + 1):
        printed = False
        for c in range(1, ncols + 1):
            cf = ws_f.cell(r, c)
            if cf.value is None or (isinstance(cf.value, str) and cf.value == ""):
                continue
            printed = True
            ref = "%s%d" % (get_column_letter(c), r)
            label = type_label(cf)
            fmt = cf.number_format
            fmt_s = "" if fmt in (None, "General") else "  fmt=%s" % short(fmt, 28)
            if label == "formula":
                cached = ws_v.cell(r, c).value
                out("  %-6s | %-20s | %s -> %s%s"
                    % (ref, label, short(formula_text(cf.value), 70),
                       short(cached, 34), fmt_s))
            else:
                out("  %-6s | %-20s | %s%s" % (ref, label, short(cf.value), fmt_s))
        if not printed:
            out("  --- row %d: entirely blank ---" % r)

    # --- per-column type census --------------------------------------------
    out("")
    out("COLUMN TYPE CENSUS  (rows 1..%d)" % min(declared_rows, args.scan_rows))
    col_types = defaultdict(Counter)
    col_samples = defaultdict(list)
    col_fmts = defaultdict(Counter)
    for row in ws_f.iter_rows(min_row=1, max_row=min(declared_rows, args.scan_rows),
                              max_col=ncols):
        for cell in row:
            label = type_label(cell)
            col_types[cell.column][label] += 1
            if label != "blank":
                col_fmts[cell.column][cell.number_format] += 1
                samples = col_samples[cell.column]
                value = formula_text(cell.value) if label == "formula" else cell.value
                rep = short(value, 46)
                if len(samples) < 5 and rep not in samples:
                    samples.append(rep)
    for c in range(1, ncols + 1):
        counts = col_types.get(c, Counter())
        non_blank = sum(v for k, v in counts.items() if k != "blank")
        if non_blank == 0:
            continue
        mix = ", ".join("%s=%d" % (k, v) for k, v in counts.most_common() if k != "blank")
        fmts = ", ".join("%s(%d)" % (short(k, 20), v) for k, v in col_fmts[c].most_common(3))
        out("  col %-4s non-blank=%-6d %s" % (get_column_letter(c), non_blank, mix))
        out("           formats : %s" % fmts)
        out("           samples : %s" % " | ".join(col_samples[c]))

    # --- blank / ragged census ---------------------------------------------
    out("")
    out("BLANK & RAGGED ROWS")
    out("  fully blank rows (in scan range): %d" % blank_rows)
    out("  first blank row numbers         : %s" % (blank_row_nums or "(none)"))
    if row_widths:
        modal_width, modal_count = row_widths.most_common(1)[0]
        ragged = sum(v for k, v in row_widths.items() if k != modal_width)
        out("  modal last-used column          : %s (%s, %d rows)"
            % (modal_width, get_column_letter(modal_width), modal_count))
        out("  rows differing from modal width : %d" % ragged)
        out("  width distribution              : %s"
            % sorted(row_widths.items(), key=lambda kv: -kv[1])[:12])

    # --- formula census -----------------------------------------------------
    shapes = Counter()
    examples = {}
    refs = Counter()
    n_formulas = 0
    for row in ws_f.iter_rows(min_row=1, max_row=min(declared_rows, args.scan_rows),
                              max_col=ncols):
        for cell in row:
            if cell.data_type != "f":
                continue
            n_formulas += 1
            text = formula_text(cell.value)
            shape = formula_shape(text)
            shapes[shape] += 1
            examples.setdefault(shape, (cell.coordinate, text,
                                        ws_v.cell(cell.row, cell.column).value))
            for sheet_name in sheets_referenced(text):
                refs[sheet_name] += 1
    out("")
    out("FORMULA CENSUS: %d formula cells, %d distinct shapes" % (n_formulas, len(shapes)))
    for shape, count in shapes.most_common(25):
        ref, text, cached = examples[shape]
        out("  x%-5d %s" % (count, short(shape, 92)))
        out("         e.g. %s: %s  -> cached %s" % (ref, short(text, 78), short(cached, 30)))
    if len(shapes) > 25:
        out("  ... %d more shapes (see --json output)" % (len(shapes) - 25))
    if refs:
        out("  cross-sheet references from this sheet:")
        for sheet_name, count in refs.most_common(20):
            out("    -> %-34s x%d" % (short(sheet_name, 34), count))

    doc_sheet.update({
        "name": name,
        "state": str(safe(lambda: ws_f.sheet_state, "?")),
        "declared_max_row": declared_rows,
        "declared_max_col": declared_cols,
        "used_last_row": last_row,
        "used_last_col": last_col,
        "merged": merged,
        "hidden_rows": hidden_rows,
        "hidden_cols": hidden_cols,
        "blank_rows": blank_rows,
        "row_width_distribution": dict(row_widths),
        "formula_count": n_formulas,
        "formula_shapes": dict(shapes.most_common()),
        "cross_sheet_refs": dict(refs),
        "tables": {k: str(v.ref) for k, v in tables.items()},
    })
    return refs


def value_fingerprint(ws_v, scan_rows, max_cols):
    """Hash of the values a human would SEE on the sheet (cached results, not
    formula text). Sheets that are duplicates or unfilled copies of a template
    collide here even when their formulas differ. Observational only."""
    parts = []
    for row in ws_v.iter_rows(min_row=1, max_row=min(ws_v.max_row, scan_rows),
                              max_col=min(ws_v.max_column, max_cols)):
        for cell in row:
            if not is_empty(cell.value):
                parts.append("%d:%d:%s" % (cell.row, cell.column, cell.value))
    if not parts:
        return "EMPTY"
    return hashlib.md5("|".join(parts).encode("utf-8", "replace")).hexdigest()[:10]


def layout_fingerprint(ws_f, rows=8, cols=30):
    """Sheets sharing this fingerprint have the same label skeleton, so one can
    stand in for many in the condensed report. Numbers and dates are ignored so
    that daily sheets with identical structure collapse together."""
    parts = []
    for r in range(1, min(rows, ws_f.max_row) + 1):
        for c in range(1, min(cols, ws_f.max_column) + 1):
            value = ws_f.cell(r, c).value
            if isinstance(value, str) and value.strip():
                parts.append("%d:%d:%s" % (r, c, value.strip().lower()[:40]))
    parts.append("cols=%d" % ws_f.max_column)
    return hashlib.md5("|".join(parts).encode("utf-8", "replace")).hexdigest()[:10]


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser(description="Profile an .xlsx workbook (read-only).")
    ap.add_argument("workbook")
    ap.add_argument("--rows", type=int, default=20)
    ap.add_argument("--max-cols", type=int, default=60)
    ap.add_argument("--scan-rows", type=int, default=5000)
    ap.add_argument("--out", default="workbook_profile.txt")
    ap.add_argument("--json", dest="json_out", default=None)
    ap.add_argument("--full", action="store_true", help="no grouping; dump every sheet")
    ap.add_argument("--reps", type=int, default=2, help="detail dumps per layout group")
    args = ap.parse_args()

    path = args.workbook
    if not os.path.exists(path):
        sys.exit("ERROR: file not found: %s" % path)
    if path.lower().endswith((".xls", ".xlsb")):
        sys.exit("ERROR: openpyxl reads .xlsx/.xlsm only. Re-save as .xlsx first.")

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    print("loading workbook (formulas)...", file=sys.stderr)
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    print("loading workbook (cached values)...", file=sys.stderr)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)

    header, index_lines, group_lines, detail = [], [], [], OrderedDict()
    doc = {"generated": _dt.datetime.now().isoformat(timespec="seconds"), "sheets": []}

    profile_workbook_meta(path, wb_f, wb_v, header.append, doc)
    profile_string_hygiene(wb_f, args, header.append, doc)

    # ---- index + fingerprints ---------------------------------------------
    index_lines.append("")
    index_lines.append("=" * 78)
    index_lines.append("SHEET INDEX  (one line per sheet -- always complete)")
    index_lines.append("=" * 78)
    index_lines.append("%-4s %-30s %-9s %-11s %-11s %5s %5s %5s %6s %-10s %-10s"
                       % ("idx", "name", "state", "declared", "used", "merg",
                          "hidR", "hidC", "formul", "layout", "content"))

    groups = defaultdict(list)
    content_groups = defaultdict(list)
    all_refs = Counter()

    for i, ws_f in enumerate(wb_f.worksheets):
        print("  profiling [%d/%d] %s" % (i + 1, len(wb_f.worksheets), ws_f.title),
              file=sys.stderr)
        ws_v = wb_v[ws_f.title]
        fp = layout_fingerprint(ws_f)
        vfp = value_fingerprint(ws_v, args.scan_rows, args.max_cols)
        groups[fp].append(ws_f.title)
        content_groups[vfp].append(ws_f.title)

        lines = []
        doc_sheet = {}
        refs = profile_sheet(ws_f, ws_v, args, lines.append, doc_sheet)
        for k, v in refs.items():
            all_refs[(ws_f.title, k)] += v
        doc_sheet["layout_fingerprint"] = fp
        doc_sheet["value_fingerprint"] = vfp
        doc["sheets"].append(doc_sheet)
        detail[ws_f.title] = lines

        index_lines.append(
            "%-4d %-30s %-9s %-11s %-11s %5d %5d %5d %6d %-10s %-10s"
            % (i, short(ws_f.title, 30).strip("'"),
               str(doc_sheet["state"])[:9],
               "%sx%s" % (doc_sheet["declared_max_row"], doc_sheet["declared_max_col"]),
               "%sx%s" % (doc_sheet["used_last_row"], doc_sheet["used_last_col"]),
               len(doc_sheet["merged"]), len(doc_sheet["hidden_rows"]),
               len(doc_sheet["hidden_cols"]), doc_sheet["formula_count"], fp, vfp))

    # ---- layout groups -----------------------------------------------------
    group_lines.append("")
    group_lines.append("=" * 78)
    group_lines.append("LAYOUT GROUPS  (sheets sharing an identical label skeleton)")
    group_lines.append("=" * 78)
    reps = []
    for fp, members in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        group_lines.append("%s  %d sheet(s)" % (fp, len(members)))
        group_lines.append("   members: %s" % ", ".join(short(m, 24) for m in members[:40]))
        if len(members) > 40:
            group_lines.append("   ... and %d more" % (len(members) - 40))
        reps.extend(members[: args.reps])
    doc["layout_groups"] = dict(groups)

    # ---- duplicate / template content -------------------------------------
    dup_lines = ["", "=" * 78,
                 "DUPLICATE CONTENT  (sheets whose visible values are identical)",
                 "=" * 78,
                 "Two sheets colliding here show the same numbers to a reader. That is",
                 "either a real duplicate, or copies of an unfilled template. Reported as",
                 "an observation -- it does not say which.", ""]
    dupes = {k: v for k, v in content_groups.items() if len(v) > 1}
    empty_sheets = content_groups.get("EMPTY", [])
    if dupes:
        for vfp, members in sorted(dupes.items(), key=lambda kv: -len(kv[1])):
            label = "  (all cells empty)" if vfp == "EMPTY" else ""
            dup_lines.append("%s  %d sheet(s)%s" % (vfp, len(members), label))
            dup_lines.append("   %s" % ", ".join(short(m, 24) for m in members[:40]))
            if len(members) > 40:
                dup_lines.append("   ... and %d more" % (len(members) - 40))
    else:
        dup_lines.append("  (no two sheets share identical visible content)")
    dup_lines.append("")
    dup_lines.append("completely empty sheets: %d %s"
                     % (len(empty_sheets),
                        [short(m, 24) for m in empty_sheets[:20]] or ""))
    doc["content_groups"] = dict(content_groups)

    # ---- cross-sheet reference map ----------------------------------------
    ref_lines = ["", "=" * 78, "CROSS-SHEET REFERENCE MAP", "=" * 78]
    if all_refs:
        for (src, dst), count in all_refs.most_common(80):
            ref_lines.append("  %-32s -> %-32s x%d" % (short(src, 32), short(dst, 32), count))
        if len(all_refs) > 80:
            ref_lines.append("  ... %d more edges" % (len(all_refs) - 80))
    else:
        ref_lines.append("  (no cross-sheet formula references found)")

    full = header + index_lines + group_lines + dup_lines + ref_lines
    for lines in detail.values():
        full.extend(lines)

    condensed = header + index_lines + group_lines + dup_lines + ref_lines
    if args.full:
        for lines in detail.values():
            condensed.extend(lines)
    else:
        condensed.append("")
        condensed.append("NOTE: showing full detail for up to %d representative sheet(s) per"
                         % args.reps)
        condensed.append("      layout group. The complete report is in %s" % args.out)
        condensed.append("      (re-run with --full to print everything to stdout).")
        for name in reps:
            condensed.extend(detail[name])

    with open(args.out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(full) + "\n")
    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as fh:
            json.dump(doc, fh, indent=2, default=str)

    print("\n".join(condensed))
    print("", file=sys.stderr)
    print("full report written to: %s (%d lines)" % (args.out, len(full)), file=sys.stderr)
    if args.json_out:
        print("json written to       : %s" % args.json_out, file=sys.stderr)


if __name__ == "__main__":
    main()
